#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 balance.csv 拆成多张按域分类的 sheet，生成 balance.xlsx。
参考 MS 按"玩家/技能/角色/敌人/Boss"分表的思路，但保留本游戏的键名体系。
"""
import os
from collections import OrderedDict

SRC = "/Users/hahaha/Documents/GitHub/vessel-gurdian/Assets/Resources/balance.csv"
DST = os.path.join(os.path.dirname(os.path.abspath(__file__)), "balance.xlsx")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def parse_value(s):
    s = s.strip()
    try:
        if "." in s or "e" in s.lower():
            return float(s)
        return int(s)
    except ValueError:
        return s


def sheet_of(cat):
    if cat == "玩家":
        return "1.玩家"
    if cat in ("攻击", "经验", "通用效果"):
        return "2.战斗基础"
    if cat in ("镜头", "关卡"):
        return "3.关卡与镜头"
    if cat in ("角色通用", "医生", "护士", "药剂师"):
        return "4.角色"
    if cat.startswith("技能-"):
        return "5.技能"
    if cat == "敌人通用" or (cat.startswith("敌人-") and cat != "敌人-Boss"):
        return "6.敌人"
    if cat in ("敌人-Boss", "Boss行为"):
        return "7.Boss"
    if cat in ("刷怪", "难度曲线", "敌人组成"):
        return "8.刷怪与难度"
    return "9.其他"


rows = []
with open(SRC, "r", encoding="utf-8-sig") as f:
    lines = f.read().split("\n")
for i, line in enumerate(lines):
    line = line.rstrip("\r")
    if i == 0 or line.strip() == "" or line.lstrip().startswith("#"):
        continue
    parts = line.split(",")
    if len(parts) < 3:
        continue
    cat = parts[0].strip()
    key = parts[1].strip()
    val = parse_value(parts[2])
    desc = ",".join(parts[3:]).strip() if len(parts) > 3 else ""
    rows.append((cat, key, val, desc))

sheets = OrderedDict()
for r in rows:
    sheets.setdefault(sheet_of(r[0]), []).append(r)

SHEET_ORDER = [
    "1.玩家", "2.战斗基础", "3.关卡与镜头", "4.角色",
    "5.技能", "6.敌人", "7.Boss", "8.刷怪与难度", "9.其他",
]

wb = Workbook()
wb.remove(wb.active)

# 说明页
intro = wb.create_sheet("0.说明")
intro_lines = [
    ["数值总表（按域拆成 8 张小表）"],
    [""],
    ["改法：只改各表“数值”列；改完运行 export_balance.py 导回游戏。"],
    ["“键名”是代码读取的钥匙，不要改。“分类”列也不要改。"],
    [""],
    ["各表内容："],
    ["1.玩家        → 玩家血量 / 移速 / 减伤"],
    ["2.战斗基础    → 攻击、经验、通用效果"],
    ["3.关卡与镜头  → 关卡时长、镜头视野"],
    ["4.角色        → 医生 / 护士 / 药剂师 + 角色通用"],
    ["5.技能        → 10 个技能的等级数值"],
    ["6.敌人        → 6 种小怪"],
    ["7.Boss        → Boss 本体与 Boss 行为"],
    ["8.刷怪与难度  → 刷怪节奏、难度曲线、敌人组成权重"],
    ["9.其他        → 未归类的（正常应为空）"],
]
for line in intro_lines:
    intro.append(line)
intro.column_dimensions["A"].width = 70
intro["A1"].font = Font(bold=True, size=13)

head_fill = PatternFill("solid", fgColor="4472C4")
head_font = Font(bold=True, color="FFFFFF", size=11)

for name in SHEET_ORDER:
    if name not in sheets or not sheets[name]:
        continue
    ws = wb.create_sheet(name)
    ws.append(["分类", "键名", "数值", "说明"])
    for r in sheets[name]:
        ws.append(list(r))
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 56
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.alignment = Alignment(horizontal="right", vertical="center")

wb.save(DST)
total = sum(len(v) for v in sheets.values())
print(f"OK 已生成 {DST}，共 {total} 行，分成 {len(sheets)} 张表：")
for name in SHEET_ORDER:
    if name in sheets:
        print(f"  {name}: {len(sheets[name])} 行")
