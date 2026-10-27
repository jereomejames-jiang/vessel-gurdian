#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 balance.xlsx（多张分表）重新拼回游戏读取的平铺 balance.csv。

用法：
  1. 用 Excel 打开 balance.xlsx，只改各表“数值”列（别改“键名”列）。
  2. 保存并关闭 Excel。
  3. 运行本脚本，会把所有分表拼成一张 balance.csv，并尝试写回游戏目录。
"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "balance.xlsx")
LOCAL_CSV = os.path.join(HERE, "balance.csv")
TARGET_CSV = "/Users/hahaha/Documents/GitHub/vessel-gurdian/Assets/Resources/balance.csv"

from openpyxl import load_workbook


def fmt_value(v):
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    return str(v)


wb = load_workbook(XLSX, data_only=True)
lines = ["分类,键名,数值,说明"]
count = 0
for sname in wb.sheetnames:
    if sname == "0.说明":
        continue
    ws = wb[sname]
    for row in ws.iter_rows(min_row=2, values_only=True):
        cat = "" if row[0] is None else str(row[0]).strip()
        key = "" if row[1] is None else str(row[1]).strip()
        if not key:
            continue
        val = fmt_value(row[2])
        desc = "" if row[3] is None else str(row[3]).strip()
        lines.append(f"{cat},{key},{val},{desc}")
        count += 1

content = "\n".join(lines) + "\n"

with open(LOCAL_CSV, "w", encoding="utf-8-sig", newline="") as f:
    f.write(content)
print(f"OK 已生成本地 {LOCAL_CSV}（{count} 行）")

try:
    with open(TARGET_CSV, "w", encoding="utf-8-sig", newline="") as f:
        f.write(content)
    print(f"OK 已写入游戏 {TARGET_CSV}")
except OSError as e:
    print(f"提示：无法直接写游戏目录（{e}）")
    print(f"请手动把 {LOCAL_CSV} 复制覆盖到：{TARGET_CSV}")
